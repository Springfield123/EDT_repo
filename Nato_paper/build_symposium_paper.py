#!/usr/bin/env python3
"""Build the IST-SET-226 short paper from the supplied legacy CSO template.

The source .doc is never modified.  The script starts from the task-local DOCX
conversion retained during template distillation, replaces the instructional
body, and preserves the template's section geometry and header/footer assets.
"""

from __future__ import annotations

import hashlib
import shutil
import sys
import csv
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
REFERENCE_DOCX = Path("/private/tmp/nato_paper_review/template/Paper_template.docx")
REFERENCE_SHA256 = "141e21d531814eaf67b458c316912374135b2fe29f69a83675de04bad406ba59"
OUTPUT_DOCX = ROOT / "Nato_paper" / "EDT_XY8_IST_SET_226_Short_Paper.docx"

FIG_MAIN = ROOT / "outputs" / "reproduced_figures" / "fig_xy8_qre_nv.png"
FIG_ABLATION = ROOT / "outputs" / "reproduced_figures" / "fig_xy8_resource_ablation.png"
PROSPECTIVE_CSV = ROOT / "data" / "processed" / "xy8_prospective_validation.csv"

FULL_TITLE = "Entanglement-Dynamic Topology for Robust Quantum Sensing in Defence Applications"
RUNNING_TITLE = "EDT CONTROL WINDOWS FOR NV-CENTRE SENSING"
REPORT_ID = "STO-MP-IST-SET-226"
PAPER_ID = "IST-SET-226"
CLASSIFICATION = "NATO UNCLASSIFIED/"
RELEASE = "Releasable to STEP, SGP."


ABSTRACT_TEXT = (
    "Quantum sensors may help defence users detect weak magnetic and radio-frequency signals, "
    "including the signatures sought in magnetic anomaly detection. Their value in practice depends "
    "on reliable operation, not only the best sensitivity achieved under ideal conditions. Longer "
    "control sequences can improve signal detection, but they also take more time and accumulate "
    "errors. We use Entanglement-Dynamic Topology (EDT) to identify a useful operating range where "
    "these benefits and costs remain balanced. We test the approach with a literature-based model "
    "of a nitrogen-vacancy centre in diamond using an established XY8 control sequence. For a "
    "200 kHz signal, the model recommends a pulse-sequence timing of 1.25 \u00b5s and 20 groups of "
    "eight pulses, for 160 pulses in total. Settings between 10 and 34 groups remain within 10% of "
    "the best predicted performance. A 16-group option uses 20% fewer pulses with less than a 1% "
    "predicted loss, offering a simpler practical choice. When pulse errors are removed from the "
    "model, longer sequences continue to improve, indicating that pulse accuracy is the main limit "
    "under the assumed conditions. The analysis converts measured hardware quality into a "
    "recommended setting and defines seven measurements that can confirm or challenge the "
    "prediction. These results are model predictions; they are not yet experimental measurements "
    "or a calibrated sensitivity in magnetic-field units."
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=50, start=70, bottom=50, end=70) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "B7B7B7")


def set_table_geometry(table, relative_widths_in: list[float], total_twips: int = 9412) -> list[int]:
    raw = [value * 1440.0 for value in relative_widths_in]
    scale = total_twips / sum(raw)
    widths = [int(round(value * scale)) for value in raw]
    widths[-1] += total_twips - sum(widths)

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    # Match the visible table border to the body start after the 70-DXA cell margin.
    tbl_ind.set(qn("w:w"), "70")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = table._tbl.find(qn("w:tblGrid"))
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
    return widths


def set_font(run, name: str, size: float | None = None, bold=None, italic=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_text(document: Document, text: str, *, style: str = "STO Text", italic=False):
    p = document.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.widow_control = True
    run = p.add_run(text)
    run.italic = italic
    return p


def add_heading(document: Document, number: str, title: str, level: int = 1):
    style = "STO First Level Heading" if level == 1 else "STO Second Level Heading"
    p = document.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = True
    p.add_run(f"{number}\t{title}")
    return p


def add_page_break(document: Document) -> None:
    p = document.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_equation(document: Document, text: str):
    p = document.add_paragraph(style="STO Text")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    set_font(run, "Cambria Math", 10.5)
    return p


def add_figure(document: Document, path: Path, width_in: float, alt_text: str):
    p = document.add_paragraph(style="STO Figure")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    inline_shape = p.add_run().add_picture(str(path), width=Inches(width_in))
    inline_shape._inline.docPr.set("descr", alt_text)
    inline_shape._inline.docPr.set("title", path.stem)
    return p


def add_caption(document: Document, text: str, *, table=False):
    style = "STO Table Caption" if table else "STO Figure Caption"
    p = document.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.add_run(text)
    return p


def format_table(table, widths: list[float], *, font_size=8.7) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    table_widths = set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for col_index, cell in enumerate(row.cells):
            cell.width = Twips(table_widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "800000")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F2F2F2")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_font(
                        run,
                        "Arial",
                        font_size,
                        bold=(row_index == 0),
                        color=(255, 255, 255) if row_index == 0 else (0, 0, 0),
                    )


def add_parameter_table(document: Document) -> None:
    add_caption(document, "Table 1: Representative model assumptions and their practical meaning.", table=True)
    rows = [
        ("Model input", "Assumption and practical meaning"),
        ("Target and timing", "200 kHz target; the primary resonant setting is 1.25 µs"),
        ("Control depth", "One XY8 block contains eight pulses; depths are tested from 1 to 128 blocks"),
        ("Optical readout", "3% contrast and 1.5 × 10¹⁰ detected photons per second; representative values"),
        ("Pulse quality", "99.7% fidelity per pulse; small errors accumulate as the sequence grows"),
        ("Coherence", "300 µs reference coherence time; control can extend coherence, but longer exposure still incurs loss"),
        ("Fixed cycle time", "7.3 µs for sensor initialization, readout, and dead time"),
        ("Operating window", "All tested choices with a predicted score no more than 10% above the best score"),
    ]
    table = document.add_table(rows=len(rows), cols=2)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    format_table(table, [1.55, 4.85], font_size=8.8)


# Short test labels for Table 2, keyed by the protocol's control_id. The numeric
# columns are read from the pipeline artifact so the table cannot drift from it.
VALIDATION_TESTS = {
    "low_depth": "Insufficient accumulation",
    "edt_lower_edge": "Lower boundary",
    "practical_choice": "Pareto compromise",
    "predicted_optimum": "Primary prediction",
    "edt_upper_edge": "Upper boundary",
    "overdriven": "Pulse-error cost",
    "off_resonance": "Timing selectivity",
}


def validation_rows() -> list[tuple[str, ...]]:
    with PROSPECTIVE_CSV.open(newline="", encoding="utf-8") as handle:
        records = list(csv.DictReader(handle))
    if len(records) != len(VALIDATION_TESTS):
        raise RuntimeError(f"Expected {len(VALIDATION_TESTS)} prospective controls, got {len(records)}.")
    rows = [("Control", "Timing (µs)", "Blocks", "Pulses", "Relative score", "Test")]
    for record in records:
        rows.append(
            (
                record["control"],
                f"{float(record['tau_s']) * 1e6:.3f}",
                record["N"],
                record["total_pi_pulses"],
                f"{float(record['predicted_relative_to_best']):.2f}",
                VALIDATION_TESTS[record["control_id"]],
            )
        )
    return rows


def add_validation_table(document: Document) -> None:
    add_caption(document, "Table 2: Prospective seven-setting validation set. Entries are predictions; no measurements have been entered.", table=True)
    rows = validation_rows()
    table = document.add_table(rows=len(rows), cols=6)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    format_table(table, [1.25, 0.74, 0.52, 0.66, 0.72, 2.42], font_size=8.2)


def add_reference(document: Document, number: int, text: str) -> None:
    p = document.add_paragraph(style="STO Reference Text")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(5)
    # The retained STO reference style supplies the bracketed sequence number.
    p.add_run(text)


def populate_document(document: Document) -> None:
    # Keep template typefaces, but tighten paragraph rhythm to satisfy the five-page main-text limit.
    document.styles["STO Text"].paragraph_format.space_after = Pt(4.5)
    document.styles["STO Abstract Text"].paragraph_format.space_after = Pt(6)
    document.styles["STO Figure Caption"].paragraph_format.space_after = Pt(6)
    document.styles["STO Table Caption"].paragraph_format.space_after = Pt(4)

    # Page 1 — title, abstract, and introduction.
    p = document.add_paragraph(style="STO Paper Title")
    p.paragraph_format.space_after = Pt(5)
    p.add_run(FULL_TITLE)

    p = document.add_paragraph(style="STO Author Name")
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    p.add_run("Hang Hu")
    r = p.add_run("1"); r.font.superscript = True
    p.add_run(", Hsu Kiang Ooi")
    r = p.add_run("1"); r.font.superscript = True
    p.add_run(", and Anguang Hu")
    r = p.add_run("2"); r.font.superscript = True

    for text in (
        "1 Digital Technologies Research Centre, Digital Technologies, National Research Council of Canada",
        "2 Suffield Research Centre, Defence Research and Development Canada, PO Box 4000 Stn Main, Medicine Hat, Alberta T1A 8K6",
        "CANADA",
    ):
        p = document.add_paragraph(style="STO Author Address")
        p.paragraph_format.space_after = Pt(0)
        p.add_run(text)

    p = document.add_paragraph(style="STO email address")
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    p.add_run("Anguang.Hu@drdc-rddc.gc.ca")

    p = document.add_paragraph(style="STO Abstract Title")
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("ABSTRACT")
    add_text(
        document,
        ABSTRACT_TEXT,
        style="STO Abstract Text",
        italic=True,
    )

    add_heading(document, "1.0", "INTRODUCTION")
    add_text(
        document,
        "NATO’s quantum strategy identifies sensing, imaging, and positioning, navigation and timing as defence-relevant application areas [1]. NV centres in diamond are attractive within this landscape because their electron spins can be initialized and read optically and can sense weak magnetic fields under ambient conditions [2–4]. Candidate uses include magnetic anomaly detection, narrowband radio-frequency sensing, and local characterization of electromagnetic signatures. For such tasks, field performance depends as much on stable control and rapid calibration as on the best value obtained in a laboratory sweep.",
    )
    add_text(
        document,
        "XY8 dynamical decoupling illustrates the engineering tension. Pulse timing selects an alternating field, while repeated eight-pulse blocks extend phase accumulation and reject unwanted spectral components [5–7]. Yet every added block consumes time and applies eight additional π rotations. Decoherence, imperfect rotations, and fixed initialization/readout time can therefore erase the nominal gain. The design question is not simply how deeply the sensor can be driven, but which finite depth remains useful for the measured hardware.",
    )
    add_text(
        document,
        "This paper directly supports the IST-SET-226 symposium aim of bridging policy, military, and specialist perspectives. It translates quantum-control calculations into two decisions that all three audiences can assess: which setting to use and which seven measurements would confirm or challenge that choice. EDT is used here as a control-window diagnostic. The result is a reproducible decision aid, not a claim of a universal topological mechanism or an operationally qualified sensor.",
    )
    add_page_break(document)

    # Page 2 — rationale and compact model specification.
    add_heading(document, "2.0", "RATIONALE")
    add_text(
        document,
        "Rather than report only one best setting, we define an EDT operating window as every tested setting predicted to perform within 10% of the best. A credible window must have clear edges: performance should worsen on both sides, and the result should remain understandable when plausible hardware assumptions change. This distinguishes a useful operating range from a single low point that may be caused by the chosen search grid.",
    )
    add_text(
        document,
        "A window is more useful than a single optimum for defence-oriented engineering. It reveals tolerance to calibration drift, supports a lower-complexity fallback, and gives test teams boundary cases that should fail for known reasons. A Pareto view also exposes the price of additional control: the operator can trade pulse count against predicted sensitivity instead of treating depth as free. In deployed or size-, weight-, and power-constrained systems, this can reduce microwave duty, calibration burden, and exposure to unmodelled errors.",
    )
    add_text(
        document,
        "The analysis asks three practical questions. Does a middle sequence depth outperform both shallower and deeper choices? Which hardware limitation creates that middle optimum? Can measured hardware quality predict the appropriate depth? Each question has a result that could disprove the model, such as the best setting occurring at a search boundary or failing to change when a modeled limitation is removed.",
    )

    add_heading(document, "3.0", "DESCRIPTION")
    add_heading(document, "3.1", "How the model ranks control choices", level=2)
    add_text(
        document,
        "The model is a ranking tool. It compares pulse timing and sequence depth using a common score, where lower is better. It deliberately does not calculate an absolute magnetic sensitivity in tesla per square-root hertz, because that would require a device-specific calibration that has not yet been measured.",
    )
    add_text(
        document,
        "Two conventions are needed to read the results. N counts complete XY8 blocks, and each block contains eight microwave pulses. The timing value shown in the paper is half the gap between pulses. At a 200 kHz target frequency, the primary resonant timing is therefore 1.25 µs. Twenty blocks contain 160 pulses and take 400 µs before initialization and readout are added.",
    )
    add_text(
        document,
        "For each candidate setting, the score rewards signal accumulation and frequency selectivity, then accounts for loss of coherence, accumulated pulse errors, optical readout quality, and fixed preparation and readout time. Published NV sensing and dynamical-decoupling work motivates these factors [2,3,5–8]. The values in Table 1 are representative assumptions rather than a calibration of a particular sensor.",
    )
    add_parameter_table(document)
    add_page_break(document)

    # Page 3 — analysis design, reproducibility, and main landscape.
    add_heading(document, "3.2", "How the prediction was tested", level=2)
    add_text(
        document,
        "The analysis systematically tests pulse timing and sequence depth, including the exact resonant timings rather than relying on a coarse numerical grid. It records the best setting and every setting within 10% of it. Extra depth values are tested near the upper edge so that the operating window has a measured boundary. Three hardware scenarios then ask whether lower, nominal, and higher overall quality lead to different recommended depths.",
    )
    add_text(
        document,
        "To identify cause, the model is rerun while removing pulse error, coherence loss, and fixed cycle time one at a time. Portability is checked at 19 target frequencies from 50 to 500 kHz. Uncertainty is assessed with 400 repeat calculations in which plausible hardware values are varied. These ranges are assumptions, not measured error bars. A final test asks which single setting performs reliably across those possible hardware states.",
    )
    add_text(
        document,
        "A separate decision map varies pulse fidelity and coherence time while keeping readout quality, photon rate, fixed timing costs, and the remaining assumptions at their nominal values. It groups the recommendation into shallow, moderate, deep, very deep, or unresolved regions. The resulting rule takes a target frequency and measured hardware quality and returns a resonant timing, recommended depth, 10% operating window, lower-pulse alternative, and warnings when the answer reaches the limits of the tested range.",
    )
    add_text(
        document,
        "All reported numbers are generated by one repository pipeline using the command python code/run_reproduction_pipeline.py --nv-only. Focused tests verify exact 200 kHz resonances, grid refinement, the Nπ = 8N convention, scenario optima, deterministic uncertainty, and agreement among CSV, JSON, and figure-source artifacts. This removes duplicated figure implementations as a source of narrative drift.",
    )
    add_figure(
        document,
        FIG_MAIN,
        6.42,
        "Three-panel modeled XY8 result. Panel A shows the normalized filter response with a 200 kilohertz target. Panel B is a heatmap of logarithmic sensitivity penalty over the half-interpulse spacing and XY8 block count, with a ten-percent contour and markers for the best, N equals 16, and N equals 64 controls. Panel C compares conservative, nominal, and optimistic optimal depths of 8, 20, and 64 blocks.",
    )
    add_caption(
        document,
        "Figure 1: Canonical XY8 sweep. (A) The pulse sequence responds most strongly at the 200 kHz target and its harmonics. (B) Yellow indicates settings closest to the best predicted performance; purple indicates a larger penalty. The black contour encloses the 10% operating window, and markers identify the best, lower-pulse, and high-depth choices. (C) Conservative, nominal, and optimistic hardware scenarios select 8, 20, and 64 blocks, respectively.",
    )
    add_page_break(document)

    # Page 4 — numerical results and mechanism.
    add_heading(document, "4.0", "RESULTS AND OPERATIONAL IMPLICATIONS")
    add_heading(document, "4.1", "Why a middle depth performs best", level=2)
    add_text(
        document,
        "At 200 kHz, the best nominal setting is a 1.25 µs timing and 20 blocks, or 160 pulses. The pulse sequence lasts 400 µs and the full measurement cycle lasts 407.3 µs. The model retains about 95.2% of the signal after coherence loss, but the accumulated pulse-error factor has fallen to 61.8%. Shallow sequences collect too little signal; very deep sequences lose too much through repeated imperfect control. Their balance creates the middle optimum. The 10% operating window spans 10 to 34 blocks. The tested values immediately outside it perform worse, so both edges are resolved rather than imposed by the search boundary.",
    )
    add_text(
        document,
        "The window has a useful engineering interpretation. N = 8 is predicted to be 17.1% worse than the optimum because it accumulates too little signal. N = 64 is 72.4% worse because 512 imperfect rotations outweigh the extra accumulation. N = 16 is Pareto-optimal: it uses 128 pulses—20% fewer than N = 20—for only a 0.97% penalty. Thus a platform with tighter microwave-duty or reliability constraints can choose N = 16 without surrendering meaningful nominal performance.",
    )

    add_heading(document, "4.2", "What creates the interior optimum", level=2)
    add_figure(
        document,
        FIG_ABLATION,
        6.38,
        "Two-panel resource ablation. Panel A compares relative sensitivity penalty against XY8 block count for nominal, no pulse error, no decoherence, and no timing overhead cases. Panel B shows best depths of 20, 128 at the tested boundary, 20, and 20, respectively.",
    )
    add_caption(
        document,
        "Figure 2: Resource ablation at 200 kHz. Removing pulse error moves the best tested depth to the N = 128 boundary; removing decoherence or timing overhead leaves N = 20. The boundary result is unresolved rather than a claimed physical optimum.",
    )
    add_text(
        document,
        "The ablation identifies the dominant cause within the model. Removing pulse errors improves the best predicted score by 71.3% and sends the optimum to the largest tested depth. Removing coherence loss improves the score by 4.77%, and removing fixed timing costs improves it by 0.90%; both leave the optimum at 20 blocks. Pulse quality is therefore the main upper-depth limit under the nominal assumptions. Hardware with different readout, environmental noise, or coherent control errors may behave differently.",
    )
    add_text(
        document,
        "Timing remains predictive across the 50–500 kHz frequency sweep: every modeled optimum follows τ = 1/(4f). The best depth rises from N = 8 at 50 kHz to N = 20 from 150 kHz onward. A 2% timing offset at 200 kHz (1.225 rather than 1.250 µs) produces a 5.39-fold predicted penalty, making resonance selectivity a particularly sharp experimental check.",
    )
    add_text(
        document,
        "Hardware uncertainty broadens the answer without changing the best fixed choice. Across 400 repeat calculations, 89.8% of the individually best depths fall inside the nominal 10-to-34-block window. Sixteen blocks is the most frequent individual answer, while 20 blocks gives the best overall performance when one setting must be chosen before the exact hardware state is known. In 95% of the sampled cases, that fixed choice is no more than 15.4% behind a setting retuned for the individual case. This difference is useful: 16 blocks is the common low-complexity answer, while 20 blocks is the stronger fixed recommendation.",
    )
    add_page_break(document)

    # Page 5 — decision rule, prospective test, limitations, conclusions.
    add_heading(document, "4.3", "From measured hardware to a testable decision", level=2)
    add_text(
        document,
        "The decision rule turns target frequency, measured pulse quality, and coherence time into a recommended pulse spacing, depth, and 10% operating window. For the nominal example of 200 kHz, 99.7% fidelity per pulse, and 300 µs coherence time, it recommends 1.25 µs and 20 blocks, or 160 pulses. It also reports the 10-to-34-block window and the 16-block lower-pulse option. Warnings appear when the hardware lies outside the tested range, the recommendation reaches a search boundary, or uncertainty spans several depth categories. Other sensor properties still require measurement; pulse quality and coherence time do not determine real performance on their own.",
    )
    add_text(
        document,
        "Table 2 defines the minimum prospective experiment. Six resonant depths test insufficient accumulation, both window edges, the practical compromise, the predicted optimum, and overdrive; one off-resonant timing control tests selectivity. N = 20 should be repeated more often to estimate variance. The protocol supports the model only if the measured optimum lies within N = 10–34, N = 8 and 64 are worse than the interior controls, the timing control is worse than resonance, and measurements are consistent with precomputed uncertainty intervals. Failure is informative because each condition challenges a different modeled resource.",
    )
    add_validation_table(document)

    add_heading(document, "5.0", "CONCLUSIONS")
    add_text(
        document,
        "The key conclusion is practical: useful quantum control occupies a finite, hardware-dependent regime. In the representative 200 kHz NV model, the resolved optimum is N = 20, while N = 16 preserves almost all predicted sensitivity with fewer pulses. The optimum moves from N = 8 to 64 as overall hardware quality improves, and the ablation shows why—pulse infidelity sets the nominal upper-depth limit. EDT is therefore most defensible here as an explicit window diagnostic and decision workflow, not as a universal mechanism.",
    )
    add_text(
        document,
        "For defence and security users, the near-term value is reduced search and clearer qualification: measure the hardware, select a bounded operating region, retain a lower-complexity fallback, and test seven controls before committing to a full sweep. The next step is experimental falsification on a calibrated NV platform, followed by a model with finite pulse widths, coherent axis and detuning errors, and measured bath spectra. Until then, all sensitivities, uncertainty intervals, and decision regions reported here remain literature-anchored projections.",
    )
    add_page_break(document)

    # Page 6 — references, excluded from the five-page main-text limit.
    add_heading(document, "6.0", "REFERENCES")
    add_reference(
        document,
        1,
        "North Atlantic Treaty Organization, “Summary of NATO’s Quantum Technologies Strategy,” Brussels, Belgium, 16 January 2024. Available: https://www.nato.int/cps/en/natohq/official_texts_221777.htm",
    )
    add_reference(
        document,
        2,
        "C. L. Degen, F. Reinhard, and P. Cappellaro, “Quantum sensing,” Reviews of Modern Physics, vol. 89, 035002, 2017, doi: 10.1103/RevModPhys.89.035002.",
    )
    add_reference(
        document,
        3,
        "J. F. Barry, J. M. Schloss, E. Bauch, M. J. Turner, C. A. Hart, L. M. Pham, and R. L. Walsworth, “Sensitivity optimization for NV-diamond magnetometry,” Reviews of Modern Physics, vol. 92, 015004, 2020, doi: 10.1103/RevModPhys.92.015004.",
    )
    add_reference(
        document,
        4,
        "J. M. Taylor, P. Cappellaro, L. Childress, L. Jiang, D. Budker, P. R. Hemmer, A. Yacoby, R. Walsworth, and M. D. Lukin, “High-sensitivity diamond magnetometer with nanoscale resolution,” Nature Physics, vol. 4, pp. 810–816, 2008, doi: 10.1038/nphys1075.",
    )
    add_reference(
        document,
        5,
        "G. de Lange, Z. H. Wang, D. Ristè, V. V. Dobrovitski, and R. Hanson, “Universal dynamical decoupling of a single solid-state spin from a spin bath,” Science, vol. 330, pp. 60–63, 2010, doi: 10.1126/science.1192739.",
    )
    add_reference(
        document,
        6,
        "F. Casola, T. van der Sar, and A. Yacoby, “Probing condensed matter physics with magnetometry based on nitrogen-vacancy centres in diamond,” Nature Reviews Materials, vol. 3, 17088, 2018, doi: 10.1038/natrevmats.2017.88.",
    )
    add_reference(
        document,
        7,
        "Ł. Cywiński, R. M. Lutchyn, C. P. Nave, and S. Das Sarma, “How to enhance dephasing time in superconducting qubits,” Physical Review B, vol. 77, 174509, 2008, doi: 10.1103/PhysRevB.77.174509.",
    )
    add_reference(
        document,
        8,
        "S. Ernst, P. J. Scheidegger, S. Diesch, and C. L. Degen, “Modeling temperature-dependent population dynamics in the excited state of the nitrogen-vacancy center in diamond,” Physical Review B, vol. 108, 085203, 2023, doi: 10.1103/PhysRevB.108.085203.",
    )


def freeze_named_complex_fields(root, names: set[str]) -> None:
    """Keep the cached result text and remove field machinery for named fields."""
    run_tag = qn("w:r")
    fld_tag = qn("w:fldChar")
    instr_tag = qn("w:instrText")
    fld_type_attr = qn("w:fldCharType")

    for paragraph in list(root.iter(qn("w:p"))):
        cursor = 0
        while True:
            children = list(paragraph)
            if cursor >= len(children):
                break
            child = children[cursor]
            begins = list(child.iter(fld_tag)) if child.tag == run_tag else []
            if not any(node.get(fld_type_attr) == "begin" for node in begins):
                cursor += 1
                continue

            end_index = None
            separate_index = None
            instruction_parts: list[str] = []
            for index in range(cursor + 1, len(children)):
                node = children[index]
                for instr in node.iter(instr_tag):
                    if instr.text:
                        instruction_parts.append(instr.text)
                for fld in node.iter(fld_tag):
                    fld_type = fld.get(fld_type_attr)
                    if fld_type == "separate" and separate_index is None:
                        separate_index = index
                    elif fld_type == "end":
                        end_index = index
                        break
                if end_index is not None:
                    break
            if end_index is None:
                cursor += 1
                continue

            instruction = " ".join(instruction_parts).strip().upper()
            field_name = instruction.split()[0] if instruction else ""
            if field_name not in names or separate_index is None:
                cursor = end_index + 1
                continue

            # Delete begin/instruction/separator/end runs and retain cached result runs.
            for index in reversed(range(cursor, end_index + 1)):
                if index <= separate_index or index == end_index:
                    paragraph.remove(children[index])
            cursor = max(0, cursor)


def stabilize_even_footer_classification(root) -> None:
    """Anchor the legacy even-page classification to a fixed page position."""
    alt_tag = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
    text_tag = qn("w:t")
    for alt in list(root.iter(alt_tag)):
        combined = "".join(node.text or "" for node in alt.iter(text_tag))
        if "NATO UNCLASSIFIED" not in combined or "STEP, SGP" not in combined:
            continue
        anchor_tag = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor"
        position_v_tag = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}positionV"
        pos_offset_tag = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}posOffset"
        wrap_tags = {
            "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}wrapTopAndBottom",
            "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}wrapSquare",
            "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}wrapTight",
            "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}wrapThrough",
        }
        wrap_none_tag = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}wrapNone"
        for anchor in alt.iter(anchor_tag):
            anchor.set("behindDoc", "0")
            anchor.set("layoutInCell", "1")
            position_v = anchor.find(position_v_tag)
            if position_v is not None:
                position_v.set("relativeFrom", "page")
                pos_offset = position_v.find(pos_offset_tag)
                if pos_offset is not None:
                    pos_offset.text = "9950000"
            for index, child in enumerate(list(anchor)):
                if child.tag in wrap_tags:
                    replacement = etree.Element(wrap_none_tag)
                    anchor.replace(child, replacement)


def add_header_footer_alt_text(root) -> None:
    """Add nonvisual descriptions to inherited header/footer drawings."""
    doc_pr_tag = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
    for doc_pr in root.iter(doc_pr_tag):
        name = (doc_pr.get("name") or "").lower()
        if "picture" in name:
            title = "NATO STO logo"
            description = "NATO Science and Technology Organization logo."
        elif "connector" in name or not name:
            title = "Decorative rule"
            description = "Decorative horizontal rule separating page content from the header or footer."
        elif "frame" in name:
            title = "Classification marking"
            description = "NATO unclassified and releasability marking."
        else:
            title = "Template page furniture"
            description = "Decorative page furniture inherited from the NATO CSO template."
        doc_pr.set("title", title)
        doc_pr.set("descr", description)


def patch_package(path: Path) -> None:
    replacements = {
        "Paper Title": RUNNING_TITLE,
        "STO-MP-xxx-###": REPORT_ID,
        "PAPER NBR": PAPER_ID,
        "NATO UNCLASSIFIED": CLASSIFICATION,
        "RELEASABLE TO ": RELEASE,
    }
    temp = path.with_suffix(".patched.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                # python-docx cannot clear lastPrinted, and the template's value
                # (2003) would otherwise ship with the submission.
                text = data.decode("utf-8")
                text = re.sub(r"<cp:lastPrinted>.*?</cp:lastPrinted>", "", text, flags=re.S)
                target.writestr(info, text.encode("utf-8"))
                continue
            if info.filename == "docProps/custom.xml":
                # The retained template carries a Grammarly document id. Keep the
                # part (and its relationship) valid but drop the inherited value.
                data = (
                    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                    b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"'
                    b' xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"/>'
                )
                target.writestr(info, data)
                continue
            if info.filename.startswith(("word/header", "word/footer")) and info.filename.endswith(".xml"):
                text = data.decode("utf-8")
                for old, new in replacements.items():
                    text = text.replace(f">{old}<", f">{new}<")
                root = etree.fromstring(text.encode("utf-8"))
                # Keep the deliberately shortened running title instead of
                # allowing LibreOffice to refresh TITLE from the long paper title.
                freeze_named_complex_fields(root, {"COMMENTS", "KEYWORDS", "TITLE"})
                if info.filename == "word/footer1.xml":
                    stabilize_even_footer_classification(root)
                add_header_footer_alt_text(root)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            target.writestr(info, data)
    temp.replace(path)


def main() -> int:
    if not REFERENCE_DOCX.exists():
        raise FileNotFoundError(f"Retained template conversion not found: {REFERENCE_DOCX}")
    actual = sha256(REFERENCE_DOCX)
    if actual != REFERENCE_SHA256:
        raise RuntimeError(f"Template hash mismatch: expected {REFERENCE_SHA256}, got {actual}")
    for figure in (FIG_MAIN, FIG_ABLATION):
        if not figure.exists():
            raise FileNotFoundError(figure)

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE_DOCX, OUTPUT_DOCX)
    document = Document(OUTPUT_DOCX)
    clear_body(document)

    props = document.core_properties
    props.title = FULL_TITLE
    props.subject = REPORT_ID
    props.author = "Hang Hu; Hsu Kiang Ooi; Anguang Hu"
    props.category = PAPER_ID
    props.comments = CLASSIFICATION.rstrip("/")
    props.keywords = RELEASE
    props.last_modified_by = "Hang Hu"
    # The retained template carries its own authoring history; clear it so the
    # submitted package does not report the template's dates as this paper's.
    props.created = datetime(2026, 9, 5, tzinfo=timezone.utc)
    props.modified = datetime(2026, 9, 5, tzinfo=timezone.utc)
    props.revision = 1

    settings = document.settings._element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")

    populate_document(document)
    document.save(OUTPUT_DOCX)
    patch_package(OUTPUT_DOCX)
    print(OUTPUT_DOCX)
    return 0


if __name__ == "__main__":
    sys.exit(main())
